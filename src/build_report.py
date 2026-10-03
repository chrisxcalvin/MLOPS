import pandas as pd
from docx import Document
from docx.shared import Inches, Pt
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP

t = pd.read_csv("report/mlflow_comparison.csv")


def label(r):
    if r.model_family == "KNN":
        extra = "" if r.n_neighbors == 5 else ", weights=distance"
        return f"n_neighbors={int(r.n_neighbors)}{extra}"
    if r.model_family == "SVM":
        return f"C={r.C:g}, kernel=rbf"
    depth = "None" if pd.isna(r.max_depth) else r.max_depth
    return f"trees={int(r.n_estimators)}, depth={depth}"


def table(d, header, rows):
    tb = d.add_table(rows=1, cols=len(header))
    tb.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        tb.rows[0].cells[i].text = h
    for row in rows:
        c = tb.add_row().cells
        for i, v in enumerate(row):
            c[i].text = str(v)
    d.add_paragraph("")


def bullets(d, items):
    for b in items:
        d.add_paragraph(b, style="List Bullet")


d = Document()
d.styles["Normal"].font.name = "Calibri"
d.styles["Normal"].font.size = Pt(11)
d.add_heading("CIA1 - MLOps Experiment: Dry Bean Classification Pipeline", 0)
d.add_paragraph("Dataset: UCI Dry Bean Dataset (13,611 samples, 16 shape features, 7 bean classes)\n"
                "Models: KNN, SVM (RBF), Random Forest\n"
                "Tools: MLflow, DVC, Kubeflow Pipelines (KFP v2), Feast, Google Vertex AI")
d.add_paragraph("Name / Reg. No.: ____________________")

# ---- Task 1
d.add_heading("Task 1: MLOps Lifecycle Analysis", 1)
d.add_picture("report/img/lifecycle.png", width=Inches(6.5))
for s, txt in [
    ("Data ingestion", "Dry Bean CSV (13,611 x 17) pulled from the UCI repository; tracked by DVC as data-v1."),
    ("Data preprocessing", "Remove 68 duplicate rows, clip outliers at 3xIQR, drop 5 features that correlate >0.98 with others "
     "(Perimeter, ConvexArea, EquivDiameter, Compactness, ShapeFactor3). Standardisation is applied inside the model pipeline (no leakage). Output tracked as data-v2."),
    ("Model training", "KNN, SVM and Random Forest, two hyper-parameter settings each, 80/20 stratified split (seed 42); every run logged to MLflow."),
    ("Model evaluation", "Train/test accuracy, weighted precision/recall/F1, confusion matrix; runs compared in MLflow; the model must pass an accuracy gate (>= 0.90)."),
    ("Deployment", "Best model (SVM, C=10) saved as a joblib/MLflow artifact and served from a Vertex AI endpoint (cloud/deploy_vertex.py); gated by the Kubeflow pipeline."),
    ("Monitoring", "Endpoint latency/error metrics, feature-skew and drift detection on the 11 input features, periodic accuracy checks on labelled samples. "
     "A drift alarm creates a new data version and re-triggers the pipeline."),
]:
    p = d.add_paragraph(style="List Bullet")
    p.add_run(s + ": ").bold = True
    p.add_run(txt)

# ---- Task 2
d.add_heading("Task 2: Experiment Tracking with MLflow", 1)
d.add_paragraph("src/train.py trains each model with two parameter settings on both dataset versions (12 runs) and logs parameters "
                "(model, hyper-parameters, data version, data md5, row/feature counts), train and test accuracy, weighted precision / recall / F1, "
                "a confusion-matrix figure and the model itself. Run `mlflow ui --backend-store-uri sqlite:///mlflow.db` to compare runs interactively "
                "(select runs, then Compare).")
table(d, ["Data", "Model", "Parameters", "Train acc", "Test acc", "Precision", "Recall", "F1"],
      [(r.data_version, r.model_family, label(r), f"{r.train_accuracy:.4f}", f"{r.test_accuracy:.4f}",
        f"{r.precision:.4f}", f"{r.recall:.4f}", f"{r.f1_score:.4f}") for _, r in t.iterrows()])
d.add_picture("report/img/model_comparison.png", width=Inches(5.8))
d.add_paragraph("Findings: all three model families land between 91% and 92.5% test accuracy. SVM (C=10) is best on both versions "
                "(0.9243 on v1, 0.9251 on v2). KNN with distance weighting and the 300-tree forest reach 100% training accuracy but only ~92% test accuracy, "
                "i.e. they overfit; SVM has the smallest train/test gap. Data v2 helped SVM slightly (+0.04 to +0.26 points) but was marginally worse for KNN and "
                "Random Forest (-0.1 to -0.5 points). With a single split these differences are small and should not be over-interpreted (cross-validation is the next step). "
                "The main value of v2 is correctness: removing duplicates avoids identical rows landing in both train and test.")

# ---- Task 3
d.add_heading("Task 3: Dataset Versioning with DVC", 1)
d.add_paragraph("Commands used (from the project root):")
for c in ["git init && dvc init",
          "dvc add data/dry_bean.csv && git add data/dry_bean.csv.dvc data/.gitignore && git commit -m \"Data v1\" && git tag data-v1",
          "python src/preprocess.py data/dry_bean.csv data/tmp.csv && mv data/tmp.csv data/dry_bean.csv",
          "dvc add data/dry_bean.csv && git commit -am \"Data v2\" && git tag data-v2",
          "dvc get . data/dry_bean.csv --rev data-v1 -o versions/dry_bean_v1.csv",
          "git diff data-v1 data-v2 -- data/dry_bean.csv.dvc"]:
    p = d.add_paragraph(style="List Bullet")
    r = p.add_run(c)
    r.font.name = "Consolas"
    r.font.size = Pt(9)
table(d, ["", "data-v1 (original)", "data-v2 (preprocessed)"], [
    ("DVC md5", "41d2e6df4a41366cfca4ad0b15b3efec", "30208868094c0850d74e44b77a138ab3"),
    ("Size", "3,827,737 bytes", "2,752,200 bytes"),
    ("Rows", "13,611", "13,543 (-68 duplicates)"),
    ("Columns", "17", "12 (5 collinear features dropped)"),
    ("Max Area", "254,616", "136,680 (outliers clipped)"),
    ("Class balance", "DERMASON 26.1% ... BOMBAY 3.8%", "DERMASON 26.2% ... BOMBAY 3.9% (unchanged)"),
    ("Best test accuracy", "0.9243 (SVM)", "0.9251 (SVM)")])
d.add_paragraph("Benefits of data versioning:")
bullets(d, [
    "Reproducibility: a Git commit plus the .dvc file pins the exact bytes (md5) used for any experiment; MLflow stores the same md5 with each run, linking model to data.",
    "Git stays small: Git stores only a tiny pointer file while the data lives in the DVC cache / remote (S3, GCS).",
    "Safe experimentation and rollback: switch between data-v1 and data-v2 with git checkout + dvc checkout.",
    "Auditability and collaboration: diffs of .dvc files show when and how a dataset changed; teammates pull identical data.",
    "Enables automation: pipelines can request a dataset by revision (the dvc_rev parameter of our Kubeflow pipeline)."])

# ---- Task 4
d.add_heading("Task 4: Workflow Automation Design (Kubeflow Pipelines)", 1)
d.add_paragraph("kubeflow/pipeline.py defines the workflow with KFP v2 and compiles to kubeflow/dry_bean_pipeline.yaml, which can be uploaded to a Kubeflow cluster "
                "or Vertex AI Pipelines (compiled successfully; not run on a cluster). kubeflow/run_local.py executes the same five stages locally as the simulated workflow "
                "(accuracy 0.9251, gate passed).")
table(d, ["Step", "Component", "What it does"], [
    ("1", "data_collection", "Fetches the dataset at a given DVC revision; emits a Dataset artifact"),
    ("2", "data_validation", "Schema, null, label-set and row-count checks; fails the run on bad data"),
    ("3", "train", "Scaling + SVM pipeline, stratified split; emits Model and test-set artifacts"),
    ("4", "evaluate", "Computes accuracy/F1, logs Metrics, returns pass/fail against threshold 0.90"),
    ("5", "deploy", "Runs only if evaluation passes (dsl.If); uploads the model to Vertex AI and deploys an endpoint")])
d.add_paragraph("Flow: data_collection -> data_validation -> train -> evaluate -> [accuracy >= 0.90] -> deploy.")
d.add_paragraph("How automation improves reproducibility and scalability:")
bullets(d, [
    "Reproducibility: each step is a containerised component with pinned inputs, parameters and a data revision, so a run can be re-executed identically; inputs/outputs are tracked as artifacts with lineage.",
    "No manual errors: validation and the quality gate run every time, so bad data or weak models never reach deployment.",
    "Scalability: each step runs in its own pod, so heavy steps can request more CPU/GPU, run in parallel (e.g. a hyper-parameter sweep) and be cached when inputs are unchanged.",
    "Continuous training: the same pipeline can be scheduled or triggered by a drift alarm or a new data version."])

# ---- Task 5
d.add_heading("Task 5: Feature Store and Cloud Platform Study", 1)
d.add_heading("Part A: Feast feature store", 2)
d.add_paragraph("Three model features were registered in a Feast repo (feature_repo/features.py): Area (pixels in the bean region), AspectRation (major / minor axis length) "
                "and ShapeFactor1. One entity (bean, key bean_id), one FeatureView (bean_features), parquet offline store and SQLite online store. Output of src/feast_demo.py:")
p = d.add_paragraph()
r = p.add_run(open("report/feast_output.txt").read())
r.font.name = "Consolas"
r.font.size = Pt(8.5)
d.add_paragraph("How Feast keeps training and inference consistent:")
bullets(d, [
    "Single definition: features are declared once in code and registered in the Feast registry; training and serving both read that definition, preventing training-serving skew.",
    "Training: get_historical_features() does a point-in-time-correct join of features to labelled entity rows, avoiding future-data leakage.",
    "Inference: materialize copies the latest values to the online store; get_online_features() returns the same features with low latency by entity key (the two outputs above match).",
    "Reuse and governance: other models can reuse the same features; the registry gives discoverability, ownership and versioning."])
d.add_heading("Part B: AWS SageMaker vs Google Cloud Vertex AI", 2)
table(d, ["Criterion", "AWS SageMaker", "Google Vertex AI"], [
    ("Services offered", "Studio, Data Wrangler, Feature Store, Pipelines, Model Registry, Clarify, Model Monitor, Ground Truth; deep S3/IAM/Lambda integration",
     "Workbench notebooks, Feature Store, Vertex Pipelines (runs KFP v2 directly), Model Registry, Model Monitoring, Model Garden; BigQuery/GCS integration"),
    ("Experiment tracking", "SageMaker Experiments, plus managed MLflow tracking", "Vertex AI Experiments / TensorBoard; MLflow can be self-hosted"),
    ("AutoML support", "SageMaker Autopilot (tabular, with explainability)", "Vertex AI AutoML (tabular, image, text, video) - broader modality coverage"),
    ("Deployment options", "Real-time, serverless, asynchronous and batch transform endpoints, multi-model endpoints, edge",
     "Online endpoints (autoscaling), batch prediction, private endpoints, custom containers"),
    ("Cost considerations", "Pay per instance-hour / per request; Savings Plans and Spot training; real-time endpoints bill while up (serverless scales to zero)",
     "Pay per node-hour for training and prediction; committed-use discounts; endpoints bill while deployed, so undeploy when idle; new accounts get free credits"),
    ("Fit for this project", "Good if data already lives in AWS", "Chosen: our KFP v2 pipeline compiles to YAML that Vertex Pipelines runs unchanged")])
d.add_paragraph("Pricing and features change often, so check the current pricing pages before launching. Note: cloud/deploy_vertex.py is written but was NOT executed "
                "(it needs a GCP project and billing). Run it before the demo if a live deployment is required.")

d.add_heading("Repository layout", 1)
bullets(d, ["data/dry_bean.csv.dvc - DVC pointer (git tags data-v1, data-v2)",
            "src/ - preprocess.py, train.py, compare.py, diagram and report builders",
            "kubeflow/ - pipeline.py, dry_bean_pipeline.yaml, run_local.py",
            "feature_repo/ - Feast definitions",
            "cloud/deploy_vertex.py - Vertex AI deployment",
            "mlflow.db - tracked experiments"])
d.save("report/MLOps_CIA1_Report.docx")

# ---------- 5-minute deck ----------
pr = Presentation()
pr.slide_width = PI(13.33)
pr.slide_height = PI(7.5)


def slide(title, items, img=None, size=22):
    s = pr.slides.add_slide(pr.slide_layouts[5])
    s.shapes.title.text = title
    s.shapes.title.text_frame.paragraphs[0].font.size = PP(32)
    wide = img is None or "life" in img
    if img:
        s.shapes.add_picture(img, PI(.8), PI(1.6), width=PI(11.7 if "life" in img else 8))
    left, top, w = (PI(.8), PI(5.0 if img else 1.8), PI(11.7)) if wide else (PI(9.0), PI(1.8), PI(4.0))
    tf = s.shapes.add_textbox(left, top, w, PI(5)).text_frame
    tf.word_wrap = True
    for i, b in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "- " + b
        p.font.size = PP(size)
        p.space_after = PP(8)


slide("MLOps Pipeline: Dry Bean Classification", ["UCI Dry Bean: 13,611 samples, 16 shape features, 7 classes",
      "Stack: DVC - MLflow - Kubeflow Pipelines - Feast - Vertex AI", "Name / Reg. No.: ____________"], size=24)
slide("1. End-to-end lifecycle", ["Ingest, preprocess, train, evaluate, deploy, monitor; drift triggers retraining"], img="report/img/lifecycle.png", size=18)
slide("2. DVC data versions", ["data-v1: original, 13,611 x 17, 3.8 MB",
      "data-v2: -68 duplicates, outliers clipped, 5 collinear features dropped: 13,543 x 12, 2.75 MB",
      "Git holds a tiny .dvc pointer; md5 logged in every MLflow run", "Demo: git diff data-v1 data-v2; dvc get --rev"])
slide("3. MLflow experiments (12 runs)", ["KNN, SVM, RF x 2 settings x 2 data versions", "Best: SVM C=10, test acc 0.925, F1 0.925",
      "KNN-15 and RF-300: 100% train acc, overfit", "v2 gain is small: honest result", "Demo: mlflow ui, compare runs"], img="report/img/model_comparison.png", size=16)
slide("4. Kubeflow pipeline", ["data_collection, data_validation, train, evaluate, then deploy if accuracy >= 0.90",
      "KFP v2 compiled to YAML; run_local.py simulates it",
      "Containers = reproducible; parallel pods + caching = scalable; quality gate protects production"])
slide("5a. Feast feature store", ["Features: Area, AspectRation, ShapeFactor1",
      "One definition serves training (point-in-time join) and inference (online store)",
      "Demo: offline and online values identical, so no training/serving skew"])
slide("5b. SageMaker vs Vertex AI", ["Both: managed pipelines, registry, monitoring, endpoints",
      "SageMaker: Autopilot, serverless/async endpoints, best if data is on AWS",
      "Vertex AI: AutoML for more modalities, runs KFP YAML natively: our choice", "Cost: pay per node-hour; undeploy idle endpoints"])
slide("Summary", ["Lifecycle automated and versioned: data, code, experiments, features, deployment",
      "Next: cross-validation, live Vertex deployment, drift monitoring"], size=24)
pr.save("report/MLOps_CIA1_Demo.pptx")
